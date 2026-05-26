"""
excel_exporter.py
-----------------
מייצא את תוצאות הניתוח לקובץ Excel מעוצב.
"""

import io
import pandas as pd
from typing import Dict
import openpyxl
from openpyxl.styles import (
    PatternFill, Font, Alignment, Border, Side
)
from openpyxl.utils import get_column_letter


COLORS = {
    "header_bg": "1A3A5C",
    "header_fg": "FFFFFF",
    "section_bg": "E8F0FE",
    "section_fg": "1A3A5C",
    "row_even": "F5F7FA",
    "row_odd": "FFFFFF",
    "warning_bg": "FFF3CD",
    "risk_high": "F8D7DA",
    "risk_mid": "FFF3CD",
    "risk_low": "D1E7DD",
    "accent": "2B7A78",
}


def _border():
    thin = Side(style='thin', color="CCCCCC")
    return Border(left=thin, right=thin, top=thin, bottom=thin)


def _header_style(ws, row, col, value, bg=None, fg=None, bold=True, size=11):
    cell = ws.cell(row=row, column=col, value=value)
    cell.font = Font(bold=bold, color=fg or COLORS["header_fg"], size=size, name="Arial")
    cell.fill = PatternFill("solid", fgColor=bg or COLORS["header_bg"])
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    cell.border = _border()
    return cell


def _data_style(ws, row, col, value, even=True, bold=False, color=None):
    cell = ws.cell(row=row, column=col, value=value)
    bg = COLORS["row_even"] if even else COLORS["row_odd"]
    cell.fill = PatternFill("solid", fgColor=color or bg)
    cell.font = Font(bold=bold, name="Arial", size=10)
    cell.alignment = Alignment(horizontal="right", vertical="center", wrap_text=True)
    cell.border = _border()
    return cell


def export_to_excel(analysis: Dict, project_name: str, city: str) -> bytes:
    """
    יצירת קובץ Excel מרובה גיליונות עם כל תוצאות הניתוח.
    מחזיר bytes של הקובץ.
    """
    wb = openpyxl.Workbook()

    # ─── גיליון 1: סיכום מנהלים ───
    ws1 = wb.active
    ws1.title = "סיכום מנהלים"
    ws1.sheet_view.rightToLeft = True
    _build_summary_sheet(ws1, analysis, project_name, city)

    # ─── גיליון 2: תמהיל דירות ───
    ws2 = wb.create_sheet("תמהיל דירות")
    ws2.sheet_view.rightToLeft = True
    _build_mix_sheet(ws2, analysis)

    # ─── גיליון 3: זכויות בנייה ───
    ws3 = wb.create_sheet("זכויות בנייה")
    ws3.sheet_view.rightToLeft = True
    _build_rights_sheet(ws3, analysis)

    # ─── גיליון 4: סיכוני תכנון ───
    ws4 = wb.create_sheet("סיכוני תכנון")
    ws4.sheet_view.rightToLeft = True
    _build_risks_sheet(ws4, analysis)

    # ─── גיליון 5: נתון גולמי JSON ───
    ws5 = wb.create_sheet("נתונים גולמיים")
    ws5.sheet_view.rightToLeft = True
    import json
    ws5["A1"] = json.dumps(analysis, ensure_ascii=False, indent=2)

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf.read()


def _build_summary_sheet(ws, analysis, project_name, city):
    ws.column_dimensions["A"].width = 30
    ws.column_dimensions["B"].width = 50
    ws.row_dimensions[1].height = 40

    _header_style(ws, 1, 1, f"דוח ניתוח תב\"ע – {project_name} | {city}", size=14)
    ws.merge_cells("A1:B1")

    meta = analysis.get("metadata", {})
    summary = analysis.get("summary", {})
    financial = analysis.get("financial_estimate", {})

    rows = [
        ("📋 פרטי תוכנית", ""),
        ("מספר תוכנית", meta.get("תוכנית_מספר", "—")),
        ("שם תוכנית", meta.get("שם_תוכנית", "—")),
        ("עיר/ישוב", meta.get("עיר_ישוב", city)),
        ("ייעוד קרקע", meta.get("ייעוד_קרקע", "—")),
        ("תאריך אישור", meta.get("תאריך_אישור", "—")),
        ("", ""),
        ("💰 הערכה כלכלית", ""),
        ("סה\"כ יחידות מוצעות", financial.get("סה_כ_יחידות_מוצעות", "—")),
        ("שטח עיקרי מנוצל מ\"ר", financial.get("שטח_עיקרי_מנוצל_מ2", "—")),
        ("אחוז ניצול זכויות", f"{financial.get('אחוז_ניצול_זכויות', '—')}%"),
        ("הכנסה משוערת (מיליון ₪)", financial.get("הכנסה_משוערת_מיליון_ש_ח", "—")),
        ("", ""),
        ("📝 סיכום מנהלים", summary.get("סיכום_מנהלים", "—")),
        ("✅ המלצה", summary.get("המלצה_עיקרית", "—")),
    ]

    for i, (k, v) in enumerate(rows):
        row = i + 2
        is_section = not v or k.startswith(("📋", "💰", "📝", "✅"))
        if k == "":
            continue

        if is_section and not v:
            cell_k = _header_style(ws, row, 1, k,
                                   bg=COLORS["section_bg"],
                                   fg=COLORS["section_fg"])
            ws.merge_cells(f"A{row}:B{row}")
        else:
            _data_style(ws, row, 1, k, even=(i % 2 == 0), bold=True)
            _data_style(ws, row, 2, str(v), even=(i % 2 == 0))
        ws.row_dimensions[row].height = 25

    # נקודות חוזק
    strengths = summary.get("נקודות_חוזק", [])
    if strengths:
        r = len(rows) + 3
        _header_style(ws, r, 1, "💪 נקודות חוזק", bg=COLORS["section_bg"], fg=COLORS["section_fg"])
        ws.merge_cells(f"A{r}:B{r}")
        for s in strengths:
            r += 1
            _data_style(ws, r, 1, "✔", even=True)
            _data_style(ws, r, 2, str(s), even=True)


def _build_mix_sheet(ws, analysis):
    ws.column_dimensions["A"].width = 18
    ws.column_dimensions["B"].width = 20
    ws.column_dimensions["C"].width = 18
    ws.column_dimensions["D"].width = 18
    ws.column_dimensions["E"].width = 22
    ws.column_dimensions["F"].width = 30

    headers = ["סוג דירה", "שטח ממוצע מ\"ר", "כמות יחידות",
               "% מהתמהיל", "שטח כולל מ\"ר", "הערה"]
    for col, h in enumerate(headers, 1):
        _header_style(ws, 1, col, h)
    ws.row_dimensions[1].height = 30

    mix = analysis.get("apartment_mix", [])
    total_units = 0
    total_area = 0

    for i, apt in enumerate(mix):
        row = i + 2
        vals = [
            apt.get("סוג_דירה", "—"),
            apt.get("שטח_ממוצע_מ2", "—"),
            apt.get("כמות_יחידות", "—"),
            f"{apt.get('אחוז_מהתמהיל', '—')}%",
            apt.get("שטח_כולל_מ2", "—"),
            apt.get("הערה", "")
        ]
        for col, v in enumerate(vals, 1):
            _data_style(ws, row, col, v, even=(i % 2 == 0))

        units = apt.get("כמות_יחידות") or 0
        area = apt.get("שטח_כולל_מ2") or 0
        total_units += units
        total_area += area

    # שורת סיכום
    summary_row = len(mix) + 2
    _header_style(ws, summary_row, 1, "סה\"כ", bg=COLORS["accent"])
    _header_style(ws, summary_row, 2, "", bg=COLORS["accent"])
    _header_style(ws, summary_row, 3, total_units, bg=COLORS["accent"])
    _header_style(ws, summary_row, 4, "100%", bg=COLORS["accent"])
    _header_style(ws, summary_row, 5, f"{total_area:,.0f}", bg=COLORS["accent"])
    _header_style(ws, summary_row, 6, "", bg=COLORS["accent"])

    # אזהרות חישוב
    warnings = analysis.get("validation_warnings", [])
    if warnings:
        r = summary_row + 2
        _header_style(ws, r, 1, "⚠️ אזהרות חישוב",
                      bg=COLORS["warning_bg"], fg="664D03")
        ws.merge_cells(f"A{r}:F{r}")
        for w in warnings:
            r += 1
            cell = ws.cell(row=r, column=1, value=f"⚠ {w.get('פירוט', '')}")
            cell.fill = PatternFill("solid", fgColor=COLORS["warning_bg"])
            cell.font = Font(name="Arial", size=10, color="664D03")
            ws.merge_cells(f"A{r}:F{r}")


def _build_rights_sheet(ws, analysis):
    ws.column_dimensions["A"].width = 32
    ws.column_dimensions["B"].width = 28

    rights = analysis.get("zoning_rights", {})
    parking = analysis.get("parking", {})
    permitted = analysis.get("permitted_uses", {})
    bonuses = analysis.get("bonuses_and_exceptions", [])

    _header_style(ws, 1, 1, "פרמטר")
    _header_style(ws, 1, 2, "ערך")
    ws.row_dimensions[1].height = 28

    sections = [
        ("📐 זכויות בנייה", {
            "שטח מגרש מ\"ר": rights.get("שטח_מגרש_מ2"),
            "אחוזי בנייה מותרים": f"{rights.get('אחוזי_בנייה_מותרים', '—')}%",
            "שטח כולל מותר מ\"ר": rights.get("שטח_כולל_מותר_מ2"),
            "שטח עיקרי מותר מ\"ר": rights.get("שטח_עיקרי_מותר_מ2"),
            "שטח שירות מותר מ\"ר": rights.get("שטח_שירות_מותר_מ2"),
            "מספר קומות מקסימלי": rights.get("מספר_קומות_מקסימלי"),
            "גובה מקסימלי מ\"": rights.get("גובה_מקסימלי_מ"),
            "קו בניין קדמי מ\"": rights.get("קו_בניין_קדמי_מ"),
            "קו בניין צדדי מ\"": rights.get("קו_בניין_צדדי_מ"),
            "אחוז כיסוי מגרש": f"{rights.get('אחוז_כיסוי_מגרש', '—')}%",
        }),
        ("🅿️ חנייה", {
            "חניות חובה למגורים (ליח\"ד)": parking.get("חניות_חובה_למגורים"),
            "חניות חובה למסחר": parking.get("חניות_חובה_למסחר"),
            "סה\"כ חניות נדרשות": parking.get("סה_כ_חניות_נדרשות"),
            "הערות": parking.get("הערות_חנייה"),
        }),
        ("🏗️ שימושים מותרים", {
            k: ("✅ כן" if v is True else "❌ לא" if v is False else str(v) if v else "—")
            for k, v in permitted.items()
        }),
    ]

    current_row = 2
    for section_title, data in sections:
        _header_style(ws, current_row, 1, section_title,
                      bg=COLORS["section_bg"], fg=COLORS["section_fg"])
        ws.merge_cells(f"A{current_row}:B{current_row}")
        ws.row_dimensions[current_row].height = 25
        current_row += 1

        for i, (k, v) in enumerate(data.items()):
            _data_style(ws, current_row, 1, k, even=(i % 2 == 0), bold=True)
            _data_style(ws, current_row, 2, str(v) if v is not None else "—",
                        even=(i % 2 == 0))
            ws.row_dimensions[current_row].height = 22
            current_row += 1
        current_row += 1

    # בונוסים
    if bonuses:
        _header_style(ws, current_row, 1, "🎁 הקלות ובונוסים",
                      bg=COLORS["section_bg"], fg=COLORS["section_fg"])
        ws.merge_cells(f"A{current_row}:B{current_row}")
        current_row += 1
        for b in bonuses:
            _data_style(ws, current_row, 1, b.get("סוג", ""), bold=True)
            _data_style(ws, current_row, 2, b.get("תיאור", ""))
            current_row += 1


def _build_risks_sheet(ws, analysis):
    ws.column_dimensions["A"].width = 16
    ws.column_dimensions["B"].width = 45
    ws.column_dimensions["C"].width = 25

    headers = ["רמת סיכון", "תיאור", "סעיף רלוונטי"]
    for col, h in enumerate(headers, 1):
        _header_style(ws, 1, col, h)

    risk_colors = {
        "גבוה": COLORS["risk_high"],
        "בינוני": COLORS["risk_mid"],
        "נמוך": COLORS["risk_low"],
    }

    risks = analysis.get("risk_flags", [])
    for i, risk in enumerate(risks):
        row = i + 2
        level = risk.get("סוג_סיכון", "—")
        color = risk_colors.get(level, COLORS["row_even"])
        for col, v in enumerate([level, risk.get("תיאור", ""), risk.get("סעיף_רלוונטי", "")], 1):
            cell = ws.cell(row=row, column=col, value=v)
            cell.fill = PatternFill("solid", fgColor=color)
            cell.font = Font(name="Arial", size=10, bold=(col == 1))
            cell.alignment = Alignment(horizontal="right", vertical="center", wrap_text=True)
            cell.border = _border()
        ws.row_dimensions[row].height = 30

    # נתוני quality
    quality = analysis.get("data_quality", {})
    if quality:
        r = len(risks) + 3
        _header_style(ws, r, 1, "איכות הנתונים",
                      bg=COLORS["section_bg"], fg=COLORS["section_fg"])
        ws.merge_cells(f"A{r}:C{r}")
        r += 1
        _data_style(ws, r, 1, "רמת ביטחון", bold=True)
        _data_style(ws, r, 2, quality.get("רמת_ביטחון", "—"))
        ws.merge_cells(f"B{r}:C{r}")
        r += 1
        _data_style(ws, r, 1, "הערות לאנליסט", bold=True)
        _data_style(ws, r, 2, quality.get("הערות_לאנליסט", "—"))
        ws.merge_cells(f"B{r}:C{r}")
